import argparse
import json
import math
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


DEFAULT_CHECKPOINT_DIR = (
    Path(__file__).resolve().parent
    / "checkpoints"
    / "qwen_3.5_2b_gemma_bridge_v3_1_smoke1000_filmcond_textsummary_instrpool"
    / "checkpoint-1000"
)


@dataclass
class ReportMetadata:
    checkpoint_dir: Path
    output_path: Path
    language: str
    generation_date: str
    checkpoint_step: int
    architecture: str
    run_name: str
    model_id: str
    summary_judgment: str


@dataclass
class ArchitectureComponentRow:
    module: str
    current_structure: str
    is_pi05_core: str
    imitation_method: str
    why: str
    risk: str


@dataclass
class MetricTable:
    title: str
    columns: list[str]
    rows: list[list[str]]
    note: str | None = None


@dataclass
class ProbeSampleRow:
    task_text: str
    frame_index: int
    wrong_instruction_shift_l2: float
    blank_image_shift_l2: float
    zero_state_shift_l2: float
    base_first_step_mae: float
    norm_saturation_fraction: float


@dataclass
class FindingRow:
    label: str
    kind: str
    detail: str


@dataclass
class Token:
    indent: int
    content: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the latest v3.1 diagnostic report as a docx.")
    parser.add_argument(
        "--checkpoint-dir",
        type=Path,
        default=DEFAULT_CHECKPOINT_DIR,
        help="Path to the target checkpoint directory.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Optional output docx path. Defaults to <checkpoint-dir>/Qwen35_v31_instrpool_diagnostic_report_zh.docx",
    )
    parser.add_argument(
        "--language",
        default="zh",
        choices=["zh"],
        help="Document language. Only zh is supported in this generator.",
    )
    return parser.parse_args()


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _strip_quotes(value: str) -> str:
    if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
        return value[1:-1]
    return value


def parse_scalar(value: str) -> Any:
    value = value.strip()
    if value == "":
        return ""
    if value in {"null", "Null", "NULL", "~"}:
        return None
    if value in {"true", "True"}:
        return True
    if value in {"false", "False"}:
        return False
    if value in {"''", '""'}:
        return ""
    value = _strip_quotes(value)
    if re.fullmatch(r"[+-]?\d+", value):
        try:
            return int(value)
        except ValueError:
            pass
    if re.fullmatch(r"[+-]?(?:\d+\.\d*|\d+|\.\d+)(?:[eE][+-]?\d+)?", value):
        try:
            return float(value)
        except ValueError:
            pass
    return value


def tokenize_yaml(text: str) -> list[Token]:
    tokens: list[Token] = []
    for raw_line in text.splitlines():
        if not raw_line.strip():
            continue
        if raw_line.lstrip().startswith("#"):
            continue
        indent = len(raw_line) - len(raw_line.lstrip(" "))
        tokens.append(Token(indent=indent, content=raw_line.strip()))
    return tokens


def parse_yaml_like(text: str) -> Any:
    tokens = tokenize_yaml(text)
    if not tokens:
        return {}
    value, index = parse_block(tokens, 0, tokens[0].indent)
    if index != len(tokens):
        raise ValueError(f"Unexpected trailing YAML tokens at index {index}")
    return value


def parse_block(tokens: list[Token], index: int, indent: int) -> tuple[Any, int]:
    if tokens[index].content.startswith("- "):
        return parse_list(tokens, index, indent)
    return parse_map(tokens, index, indent)


def parse_map(tokens: list[Token], index: int, indent: int) -> tuple[dict[str, Any], int]:
    result: dict[str, Any] = {}
    while index < len(tokens):
        token = tokens[index]
        if token.indent < indent:
            break
        if token.indent != indent or token.content.startswith("- "):
            break
        key, separator, rest = token.content.partition(":")
        if not separator:
            raise ValueError(f"Invalid YAML mapping line: {token.content}")
        key = key.strip()
        rest = rest.strip()
        index += 1
        if rest:
            result[key] = parse_scalar(rest)
            continue
        if index >= len(tokens):
            result[key] = None
            continue
        next_token = tokens[index]
        if next_token.indent < indent:
            result[key] = None
            continue
        if next_token.content.startswith("- "):
            child_indent = next_token.indent
            result[key], index = parse_list(tokens, index, child_indent)
            continue
        if next_token.indent > indent:
            result[key], index = parse_block(tokens, index, next_token.indent)
            continue
        result[key] = None
    return result, index


def parse_list(tokens: list[Token], index: int, indent: int) -> tuple[list[Any], int]:
    items: list[Any] = []
    while index < len(tokens):
        token = tokens[index]
        if token.indent < indent:
            break
        if token.indent != indent or not token.content.startswith("- "):
            break
        rest = token.content[2:].strip()
        index += 1
        if not rest:
            if index < len(tokens) and (
                tokens[index].indent > indent or tokens[index].content.startswith("- ")
            ):
                child_indent = tokens[index].indent
                child_value, index = parse_block(tokens, index, child_indent)
                items.append(child_value)
            else:
                items.append(None)
            continue
        if ":" in rest and not rest.startswith(("'", '"')):
            key, separator, value = rest.partition(":")
            if separator:
                mapping: dict[str, Any] = {key.strip(): parse_scalar(value.strip()) if value.strip() else None}
                if index < len(tokens) and tokens[index].indent > indent:
                    child_value, index = parse_block(tokens, index, tokens[index].indent)
                    if isinstance(child_value, dict):
                        mapping.update(child_value)
                items.append(mapping)
                continue
        items.append(parse_scalar(rest))
    return items, index


def load_yaml(path: Path) -> dict[str, Any]:
    parsed = parse_yaml_like(path.read_text(encoding="utf-8"))
    if not isinstance(parsed, dict):
        raise ValueError(f"Expected YAML mapping at top level in {path}")
    return parsed


def require_file(path: Path) -> Path:
    if not path.exists():
        raise FileNotFoundError(f"Required artifact is missing: {path}")
    return path


def format_float(value: Any, digits: int = 4) -> str:
    if value is None:
        return "N/A"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return f"{value:.{digits}f}"
    return str(value)


def format_percent(value: float, digits: int = 2) -> str:
    return f"{value * 100:.{digits}f}%"


def format_bool(value: Any) -> str:
    return "是" if bool(value) else "否"


def vector_l2_norm(values: Any) -> float:
    if not isinstance(values, list):
        return 0.0
    total = 0.0
    for item in values:
        try:
            number = float(item)
        except (TypeError, ValueError):
            continue
        total += number * number
    return math.sqrt(total)


def add_run_with_font(run, font_name: str, font_size: int, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(font_size)
    run.bold = bold


def set_document_defaults(document: Document) -> None:
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal_style.font.size = Pt(10.5)

    for style_name, font_size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(font_size)

    if "DiagTable" not in document.styles:
        style = document.styles.add_style("DiagTable", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(9)


def add_paragraph_text(document: Document, text: str, bold: bool = False, align: Any = None) -> None:
    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    add_run_with_font(run, "Microsoft YaHei", 10.5, bold=bold)


def add_labeled_paragraph(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"[{label}] ")
    add_run_with_font(label_run, "Microsoft YaHei", 10.5, bold=True)
    text_run = paragraph.add_run(text)
    add_run_with_font(text_run, "Microsoft YaHei", 10.5)


def fill_cell(cell, text: str, font_size: int = 9, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    add_run_with_font(run, "Microsoft YaHei", font_size, bold=bold)


def add_table(document: Document, table_data: MetricTable, font_size: int = 9) -> None:
    if table_data.title:
        document.add_paragraph(table_data.title, style="Heading 2")
    table = document.add_table(rows=1, cols=len(table_data.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    for idx, column in enumerate(table_data.columns):
        fill_cell(header_cells[idx], column, font_size=font_size, bold=True)
    for row in table_data.rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            fill_cell(cells[idx], value, font_size=font_size)
    if table_data.note:
        add_labeled_paragraph(document, "事实", table_data.note)


def collect_code_evidence(repo_root: Path) -> dict[str, bool]:
    backbone_text = (repo_root / "model" / "qwen_backbone_adapter.py").read_text(encoding="utf-8")
    bridge_text = (repo_root / "model" / "qwen_gemma_bridge_action_head.py").read_text(encoding="utf-8")
    pi05_text = (repo_root / "model" / "qwen35_pi05_action_head.py").read_text(encoding="utf-8")
    return {
        "adapter_api_present": all(
            snippet in backbone_text
            for snippet in ["def encode_prefix", "def get_full_attention_taps", "def get_prefix_memory"]
        ),
        "text_instruction_signals_present": all(
            snippet in bridge_text
            for snippet in [
                "memory_summary_proj",
                "text_summary_proj",
                "instruction_summary_proj",
                "use_action_input_conditioning",
            ]
        ),
        "gemma_cross_attention_present": all(
            snippet in bridge_text for snippet in ["class GemmaExpertBlock", "self.cross_attn", "class GemmaActionExpert"]
        ),
        "pi05_dual_stream_present": all(
            snippet in pi05_text
            for snippet in ["class QwenPI05AdaRMSNorm", "def build_prefix_cache", "def _run_dual_stream_layers"]
        ),
    }


def load_artifacts(checkpoint_dir: Path) -> dict[str, Any]:
    config_path = require_file(checkpoint_dir / "config.yaml")
    eval_metrics_path = require_file(checkpoint_dir / "eval_metrics.json")
    open_loop_path = require_file(checkpoint_dir / "open_loop_eval.json")
    vuln_path = require_file(checkpoint_dir / "probe_model_vulnerabilities_libero_spatial.json")
    semantics_path = require_file(checkpoint_dir / "probe_libero_action_semantics_task0.json")
    closed_loop_path = require_file(checkpoint_dir / "eval_libero_official_libero_spatial_tasks_0_2_10eps.json")
    preprocessor_path = require_file(checkpoint_dir / "policy_preprocessor.json")
    postprocessor_path = require_file(checkpoint_dir / "policy_postprocessor.json")
    return {
        "config": load_yaml(config_path),
        "eval_metrics": load_json(eval_metrics_path),
        "open_loop": load_json(open_loop_path),
        "vulnerability": load_json(vuln_path),
        "semantics": load_json(semantics_path),
        "closed_loop": load_json(closed_loop_path),
        "preprocessor": load_json(preprocessor_path),
        "postprocessor": load_json(postprocessor_path),
    }


def build_metadata(checkpoint_dir: Path, output_path: Path, language: str, artifacts: dict[str, Any]) -> ReportMetadata:
    config = artifacts["config"]
    run_name = str(config["training"]["wandb_run_name"])
    architecture = str(config["model"]["architecture"])
    model_id = str(config["model"]["vlm_model_id"])
    checkpoint_step = int(artifacts["eval_metrics"].get("step", checkpoint_dir.name.removeprefix("checkpoint-")))
    return ReportMetadata(
        checkpoint_dir=checkpoint_dir,
        output_path=output_path,
        language=language,
        generation_date=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        checkpoint_step=checkpoint_step,
        architecture=architecture,
        run_name=run_name,
        model_id=model_id,
        summary_judgment="当前 v3.1 已解决“指令几乎不进入 expert”的问题，但仍未达到 pi0.5 式 token-level control。",
    )


def build_core_config_rows(config: dict[str, Any], postprocessor: dict[str, Any]) -> list[list[str]]:
    head = config["action_head"]
    model = config["model"]
    runtime_bridge = postprocessor.get("runtime", {}).get("bridge", {})
    hidden_dim = int(head["hidden_dim"])
    num_heads = int(head["num_heads"])
    rows = [
        ["架构标识", str(model["architecture"])],
        ["action head 类型", str(head["head_type"])],
        ["action expert 变体", str(head["action_expert_variant"])],
        ["hidden_dim", str(hidden_dim)],
        ["num_heads", str(num_heads)],
        ["num_kv_heads", str(head["num_kv_heads"])],
        ["head_dim", str(hidden_dim // max(num_heads, 1))],
        ["num_layers", str(head["num_layers"])],
        ["mlp_dim", str(head["mlp_dim"])],
        ["vlm_hidden_dim", str(head["vlm_hidden_dim"])],
        ["bridge_out_dim", str(head["bridge_out_dim"])],
        ["tap_strategy", str(head["tap_strategy"])],
        ["conditioning_mode", str(runtime_bridge.get("conditioning_mode", head["conditioning_mode"]))],
        ["freeze_vlm", str(model["freeze_vlm"]).lower()],
        ["stop_gradient_backbone", str(head["stop_gradient_backbone"]).lower()],
        ["bridge_norm_type", str(head["bridge_norm_type"])],
        ["bridge_gate_bias", format_float(head["bridge_gate_bias"])],
        ["conditioning_gate_bias", format_float(head["conditioning_gate_bias"])],
        ["memory_summary_gain_init", format_float(head["memory_summary_gain_init"])],
        ["text_summary_gain_init", format_float(head["text_summary_gain_init"])],
        ["instruction_summary_gain_init", format_float(head["instruction_summary_gain_init"])],
        ["action_input_gain_init", format_float(head["action_input_gain_init"])],
        ["memory_norm_ratio_limit", format_float(head["memory_norm_ratio_limit"])],
        ["action_dim", str(head["action_dim"])],
        ["state_dim", str(head["state_dim"])],
        ["chunk_size", str(head["chunk_size"])],
        ["action_horizon", str(head["action_horizon"])],
        ["num_inference_steps", str(head["num_inference_steps"])],
    ]
    return rows


def build_architecture_rows() -> list[ArchitectureComponentRow]:
    return [
        ArchitectureComponentRow(
            module="action chunk / flow matching / reverse-time Euler",
            current_structure="用 50-step action chunk 预测整段动作；训练时对动作与噪声做 flow matching；推理时从 t=1 到 t=0 做 Euler 反向积分。",
            is_pi05_core="是",
            imitation_method="不适用。",
            why="这是当前仓库里最接近 pi0.5 控制协议的骨架，决定了训练目标和采样方式。",
            risk="这部分协议本身并不是主要疑点；更可能的问题出在条件如何进入 expert，而不是 flow matching 是否存在。",
        ),
        ArchitectureComponentRow(
            module="timestep conditioning",
            current_structure="用 sinusoidal timestep embedding 再接两层 MLP，得到全局时间条件向量 expert_cond 的一部分。",
            is_pi05_core="是",
            imitation_method="不适用。",
            why="保证 denoising / flow 过程在各步具有可辨识时间语义。",
            risk="时间条件本身已具备，当前没有证据表明 closed-loop 弱主要是 timestep 路径失效。",
        ),
        ArchitectureComponentRow(
            module="state prompt 路径",
            current_structure="保留 pi0.5-style 的离散 state prompt，把量化后的状态文本拼进 Qwen 指令前缀。",
            is_pi05_core="部分是",
            imitation_method="延续现有 prompt-state 协议，让 backbone 仍能在前缀里看到状态描述。",
            why="避免完全丢掉已有 Qwen VLA 文本输入接口，同时与旧数据流兼容。",
            risk="文本 state 是粗粒度离散条件，表达力不如连续状态直通 expert。",
        ),
        ArchitectureComponentRow(
            module="continuous normalized state 路径",
            current_structure="把归一化后的连续状态送入独立 MLP，再加到 expert_cond。",
            is_pi05_core="是",
            imitation_method="不适用。",
            why="让状态不仅存在于 prompt，也直接作为 expert 的连续控制信号。",
            risk="如果全局条件只以 pooled vector 注入，状态虽然进入了 expert，但可能仍不足以精细塑造轨迹。",
        ),
        ArchitectureComponentRow(
            module="Qwen full forward prefix encoding",
            current_structure="Qwen 跑完整 multimodal prefix forward；图像和文本都进入 prefix，不是只抽图像编码器特征。",
            is_pi05_core="否",
            imitation_method="用完整 prefix 编码去接近 pi0.5 中 prefix 对后续动作流的上下文承载作用。",
            why="这样可以保留完整指令与视觉上下文，再从中抽取给 action expert 消费的 memory。",
            risk="Qwen 在前缀内部已经完成了自己的混合与压缩，后续 expert 只能读到被抽取后的表示，而不是层内共同演化的 token 流。",
        ),
        ArchitectureComponentRow(
            module="last full-attn tap memory",
            current_structure="只取最后一个 full-attention layer 的 hidden state 作为对外暴露的 prefix memory。",
            is_pi05_core="否",
            imitation_method="把最后层 full-attn hidden 当成 prefix 最终摘要记忆，近似替代 pi0.5 的 prefix/suffix 共享层交互。",
            why="实现简单，显存可控，也能与 stop-gradient stage-1 训练策略兼容。",
            risk="这是明显的信息瓶颈：中间层中更干净的语言结构和细粒度控制线索可能在最后一层已被压平或改写。",
        ),
        ArchitectureComponentRow(
            module="bridge",
            current_structure="prefix memory 先过 LayerNorm -> Linear(2048->1024) -> learned gate/scale，再映到 Gemma expert 空间。",
            is_pi05_core="否",
            imitation_method="显式投影桥替代 pi0.5 那种层内共享注意力空间，让 Qwen hidden 至少能被 expert 消费。",
            why="Qwen backbone hidden_dim 与 Gemma expert hidden_dim 不同，必须有一个适配层。",
            risk="bridge 是单点瓶颈；如果投影过粗，expert 读到的是被压缩后的 memory，而不是 token-level 的原生交互信号。",
        ),
        ArchitectureComponentRow(
            module="memory summary / text summary / instruction summary",
            current_structure="bridged memory 做全 token masked mean 得到 memory_summary；对 text token masked mean 得到 text_summary；对最终层文本做 pooled summary 得到 instruction_summary。",
            is_pi05_core="否",
            imitation_method="这是专门为模拟 pi0.5 更强的 instruction-to-action 条件路径而补出来的替代件。",
            why="cross-attention only 的版本里文本信号会被视觉 token 稀释，所以需要显式地把文本摘要再送入 expert_cond。",
            risk="summary 条件是全局 pooled vector，能提升敏感度，但天然比 token-level 传播更粗，容易变成 steering 而非精确控制。",
        ),
        ArchitectureComponentRow(
            module="action-token global conditioning",
            current_structure="action encoder 输出的 token 还会额外加上 action_input_gain * expert_cond，让每个 action token 一开始就带全局条件偏置。",
            is_pi05_core="否",
            imitation_method="用输入层的全局偏置去模拟 pi0.5 suffix stream 从起点就被条件调制的效果。",
            why="避免 expert 过度依赖后续 cross-attn 才感知条件，增强训练初期的条件可见度。",
            risk="这会提升条件存在感，但也可能把条件变成粗暴偏置，未必等同于层内细粒度控制。",
        ),
        ArchitectureComponentRow(
            module="Gemma self-attn / cross-attn / MLP expert",
            current_structure="独立 Gemma-style expert stack；每层按 self-attn -> cross-attn(memory) -> MLP 顺序更新 action token。",
            is_pi05_core="否",
            imitation_method="用 Gemma 风格 expert 去扮演 pi0.5 的 action suffix expert，但保留独立 action stream 的思想。",
            why="满足你要的“Qwen backbone 接 Gemma action experts”这条新架构目标。",
            risk="operator family 已与 pi0.5-style Qwen suffix blocks 分离，功能可以近似，但归纳偏置和层间动力学并不等价。",
        ),
        ArchitectureComponentRow(
            module="per-layer AdaRMS / gated residual",
            current_structure="Gemma expert 的 self-attn、cross-attn、MLP 前都做 AdaRMS/FiLM 调制，gate bias 初始为非零。",
            is_pi05_core="是",
            imitation_method="不适用。",
            why="这部分是直接沿着 pi0.5-style 的 per-layer conditioning 思路来做的，而不是只在输入层加 condition。",
            risk="虽然形式上接近 pi0.5-style，但条件来源仍主要是 summary 向量，所以“有分层调制”不等于“有 token-level control”。",
        ),
    ]


def build_structure_metric_table(metadata: ReportMetadata, artifacts: dict[str, Any], code_evidence: dict[str, bool]) -> MetricTable:
    config = artifacts["config"]
    head = config["action_head"]
    model = config["model"]
    eval_metrics = artifacts["eval_metrics"]
    rows = [
        ["checkpoint_step", str(metadata.checkpoint_step), "eval_metrics.json"],
        ["smoke_max_steps", str(config["training"]["max_steps"]), "config.yaml"],
        ["freeze_vlm", str(model["freeze_vlm"]).lower(), "config.yaml"],
        ["stop_gradient_backbone", str(head["stop_gradient_backbone"]).lower(), "config.yaml"],
        ["tap_strategy", str(head["tap_strategy"]), "config.yaml"],
        ["conditioning_mode", str(head["conditioning_mode"]), "config.yaml"],
        ["memory_norm_ratio_limit", format_float(head["memory_norm_ratio_limit"]), "config.yaml"],
        ["eval_action_mse_at_checkpoint", format_float(eval_metrics.get("action_mse")), "eval_metrics.json"],
        ["adapter_api_present", format_bool(code_evidence["adapter_api_present"]), "static code"],
        ["gemma_cross_attention_present", format_bool(code_evidence["gemma_cross_attention_present"]), "static code"],
        ["text_instruction_signals_present", format_bool(code_evidence["text_instruction_signals_present"]), "static code"],
        ["pi05_dual_stream_reference_present", format_bool(code_evidence["pi05_dual_stream_present"]), "static code"],
    ]
    note = (
        "本组只报告 checkpoint artifact 和静态代码结构中可以直接验证的结构/梯度约束。"
        "单测运行时产生的 bridge_off_parity、显式梯度计数等标量没有随 checkpoint 持久化，因此这里不伪造不存在的数值。"
    )
    return MetricTable(
        title="第四章 诊断参数与行为证据：4.1 unit/smoke 结构与梯度指标",
        columns=["指标", "当前值", "来源"],
        rows=rows,
        note=note,
    )


def build_open_loop_table(open_loop: dict[str, Any]) -> MetricTable:
    rows = [
        ["evaluated_examples", str(open_loop["evaluated_examples"])],
        ["num_inference_steps", str(open_loop["num_inference_steps"])],
        ["raw_action_mae", format_float(open_loop["raw_action_mae"])],
        ["raw_action_mse", format_float(open_loop["raw_action_mse"])],
        ["first_step_mae", format_float(open_loop["first_step_mae"])],
        ["first_step_mse", format_float(open_loop["first_step_mse"])],
        ["non_gripper_mae", format_float(open_loop["non_gripper_mae"])],
        ["non_gripper_mse", format_float(open_loop["non_gripper_mse"])],
        ["gripper_accuracy", format_percent(float(open_loop["gripper_accuracy"]))],
        ["elapsed_min", format_float(open_loop["elapsed_min"])],
    ]
    return MetricTable(
        title="第四章 诊断参数与行为证据：4.2 open-loop 指标",
        columns=["指标", "当前值"],
        rows=rows,
    )


def build_vulnerability_table(vulnerability: dict[str, Any]) -> MetricTable:
    rows = [
        ["base_first_step_mae", format_float(vulnerability["base_first_step_mae"])],
        ["wrong_instruction_shift_l2", format_float(vulnerability["wrong_instruction_shift_l2"])],
        ["blank_image_shift_l2", format_float(vulnerability["blank_image_shift_l2"])],
        ["zero_state_shift_l2", format_float(vulnerability["zero_state_shift_l2"])],
        ["wrong_instruction_mae", format_float(vulnerability["wrong_instruction_mae"])],
        ["blank_image_mae", format_float(vulnerability["blank_image_mae"])],
        ["zero_state_mae", format_float(vulnerability["zero_state_mae"])],
        ["mean_norm_saturation_fraction", format_float(vulnerability["mean_norm_saturation_fraction"])],
        ["per_sample_count", str(len(vulnerability.get("per_sample", [])))],
    ]
    return MetricTable(
        title="第四章 诊断参数与行为证据：4.3 vulnerability probe 指标",
        columns=["指标", "当前值"],
        rows=rows,
    )


def build_closed_loop_table(closed_loop: dict[str, Any]) -> MetricTable:
    suite = next(iter(closed_loop.values()))
    rows = []
    for task_name, task in suite["tasks"].items():
        rows.append(
            [
                str(task["task_id"]),
                task_name,
                str(task["successes"]),
                str(task["episodes"]),
                format_percent(float(task["success_rate"])),
            ]
        )
    rows.append(["-", "average_success_rate", "-", "-", format_percent(float(suite["average_success_rate"]))])
    return MetricTable(
        title="第四章 诊断参数与行为证据：4.4 closed-loop 2x10 official 子集指标",
        columns=["task_id", "task", "successes", "episodes", "success_rate"],
        rows=rows,
    )


def build_findings(artifacts: dict[str, Any], code_evidence: dict[str, bool]) -> list[FindingRow]:
    open_loop = artifacts["open_loop"]
    vulnerability = artifacts["vulnerability"]
    closed_loop = next(iter(artifacts["closed_loop"].values()))
    findings = [
        FindingRow(
            label="事实 1",
            kind="事实",
            detail=(
                f"open-loop 的 raw_action_mae={format_float(open_loop['raw_action_mae'])}，"
                f"gripper_accuracy={format_percent(float(open_loop['gripper_accuracy']))}，说明 smoke1000 checkpoint"
                " 已能稳定输出有意义的动作分布，并非数值崩坏。"
            ),
        ),
        FindingRow(
            label="事实 2",
            kind="事实",
            detail=(
                f"wrong_instruction_shift_l2={format_float(vulnerability['wrong_instruction_shift_l2'])}，"
                f"zero_state_shift_l2={format_float(vulnerability['zero_state_shift_l2'])}，"
                "这说明指令和状态都已经能改变 expert 输出，不再是“条件根本进不去”的状态。"
            ),
        ),
        FindingRow(
            label="事实 3",
            kind="事实",
            detail=(
                f"blank_image_shift_l2={format_float(vulnerability['blank_image_shift_l2'])} 高于 "
                f"wrong_instruction_shift_l2={format_float(vulnerability['wrong_instruction_shift_l2'])}，"
                "说明视觉仍然是当前模型里更强的主导条件源。"
            ),
        ),
        FindingRow(
            label="事实 4",
            kind="事实",
            detail=(
                f"closed-loop 2x10 official 子集平均成功率只有 {format_percent(float(closed_loop['average_success_rate']))}，"
                "其中 task 0 为 1/10，task 2 为 0/10。运行时是稳定的，但策略还不够可靠。"
            ),
        ),
        FindingRow(
            label="事实 5",
            kind="事实",
            detail=(
                "当前 v3.1 的 runtime contract 明确使用 last full-attn tap、cross-attention bridge、"
                "per_layer_film、freeze_vlm=true、stop_gradient_backbone=true。"
            ),
        ),
        FindingRow(
            label="事实 6",
            kind="事实",
            detail=(
                "静态代码同时保留了 pi0.5-style dual-stream reference 和当前 Gemma bridge expert，"
                "这意味着我们可以明确看出：现在的问题不是仓库里没有 pi0.5-style 参考，而是 v3.1 为了接 Gemma 有意换了交互拓扑。"
                if code_evidence["pi05_dual_stream_present"]
                else "静态代码证据不完整，无法完整重建 pi0.5-style dual-stream 参考。"
            ),
        ),
        FindingRow(
            label="推断 1",
            kind="推断",
            detail=(
                "最大的不对齐更像是交互拓扑不一致，而不是 expert 容量不足。因为条件已经能改变输出，但 closed-loop 仍很弱，"
                "说明问题更可能出在“条件如何进入动作流”而不是“条件有没有进入”。"
            ),
        ),
        FindingRow(
            label="推断 2",
            kind="推断",
            detail=(
                "summary-based conditioning 是有效的补救，但它的表达粒度比 token-level 传播更粗。"
                "这能解释为什么 wrong-instruction sensitivity 已经被拉起来，但动作执行仍不够稳。"
            ),
        ),
        FindingRow(
            label="推断 3",
            kind="推断",
            detail=(
                "Gemma operator family 与仓库中的 pi0.5-style Qwen suffix blocks 并不相同。"
                "如果 closed-loop 主要受结构归纳偏置影响，那么仅靠 gain 调大不一定能补齐。"
            ),
        ),
        FindingRow(
            label="推断 4",
            kind="推断",
            detail=(
                "只取 last full-attn tap 很可能把更早层里对语言更干净的结构信息压成了最终层的混合表示，"
                "这是当前信息瓶颈里最容易被验证、也最值得优先怀疑的一项。"
            ),
        ),
    ]
    return findings


def build_suspicion_rows() -> list[list[str]]:
    return [
        [
            "1",
            "交互拓扑不一致",
            "pi0.5-style 更接近 prefix/suffix token 在层内共同演化；当前 v3.1 是 Qwen 编完 prefix 后，再让 Gemma 通过 external memory + cross-attn 去读取。",
            "因为 wrong_instruction_shift 和 zero_state_shift 已不低，说明条件存在；若只是条件缺失，不会先看到 sensitivity 恢复再看到 closed-loop 依然弱。",
            "这是最像主因的结构性差异。",
        ],
        [
            "2",
            "summary 条件过粗",
            "memory/text/instruction summary 都是 pooled vector，它们能拉高 probe，但无法天然承载 token-level 对齐信息。",
            "如果只是训练步数太少，通常会看到整体 sensitivity 与行为一起都偏弱；现在则是 sensitivity 提升明显、行为提升有限。",
            "很可能造成 steering 强、精确控制弱。",
        ],
        [
            "3",
            "Gemma/Qwen operator family 不一致",
            "当前 expert 的 block family 是 Gemma 风格；仓库里的 pi0.5-style 参考仍是 Qwen-like dual-stream suffix expert。",
            "如果只是 bridge 数值尺度不对，通常先会出现 NaN、加载异常或完全无响应；目前运行稳定，因此更像是归纳偏置差异。",
            "属于中高优先级嫌疑点。",
        ],
        [
            "4",
            "只取 last tap 的信息瓶颈",
            "最后层表示已经混合过视觉和文本，也可能丢掉中间层更结构化的语言控制信号。",
            "这比“没有文本路径”更像问题，因为文本路径现在已经被 summary 和 instruction pooled summary 明确补上了。",
            "很适合做低成本验证实验。",
        ],
        [
            "5",
            "条件增强后的 steering 过粗",
            "text/instruction/action-input gain 的设计让条件更容易被 expert 看见，但也可能把条件变成粗偏置，而不是精细 shaping。",
            "如果主要问题是数值不稳定，会先看到训练或推理异常；现在更多是能跑、也敏感，但执行不稳。",
            "需要通过更细粒度注入方式来区分。",
        ],
        [
            "6",
            "瓶颈已经从“条件缺失”转向“条件进入方式不够像 pi0.5”",
            "当前版本最重要的变化是指令已能进入 expert，但进入的载体主要是 summary 向量和 external memory，而不是 token-level dual-stream coupling。",
            "这比“数据集坏了”更像问题，因为 open-loop 预测和 probe 已证明模型并非完全失学。",
            "这是对当前状态的总体诊断结论。",
        ],
    ]


def build_next_steps() -> list[str]:
    return [
        "优先验证 multi-tap bridge：把最后若干个 full-attn tap 做 learned mix 或分层桥接。如果 closed-loop 提升明显，说明 last-tap 信息瓶颈是真问题。",
        "优先验证 token-level text injection：减少 summary-only 依赖，让 instruction token 或 text-only memory 更直接参与 expert 层内交互。如果 success rate 提升多于 probe 指标提升，说明 summary 条件确实过粗。",
        "验证 Qwen-like suffix expert 或 hybrid block：在保持 Qwen backbone 的前提下，让 action expert 更接近 pi0.5-style operator family。如果这一路改善明显，说明 Gemma/Qwen family mismatch 是关键结构误差。",
        "验证 stage-2 弱解冻：仅对 Qwen 晚期 full-attn 层或 bridge 相邻层做小幅可训练适配。如果 sensitivity 与 closed-loop 同时改善，说明 frozen backbone 对新 bridge 协议过于僵硬。",
    ]


def build_probe_sample_rows(vulnerability: dict[str, Any]) -> list[ProbeSampleRow]:
    rows: list[ProbeSampleRow] = []
    for sample in vulnerability.get("per_sample", []):
        rows.append(
            ProbeSampleRow(
                task_text=str(sample.get("task_text", "")),
                frame_index=int(sample.get("frame_index", 0)),
                wrong_instruction_shift_l2=float(sample.get("wrong_instruction_shift_l2", 0.0)),
                blank_image_shift_l2=float(sample.get("blank_image_shift_l2", 0.0)),
                zero_state_shift_l2=float(sample.get("zero_state_shift_l2", 0.0)),
                base_first_step_mae=float(sample.get("base_first_step_mae", 0.0)),
                norm_saturation_fraction=float(sample.get("norm_saturation_fraction", 0.0)),
            )
        )
    return rows


def build_action_semantics_rows(semantics: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for probe in semantics.get("probes", []):
        delta = probe.get("delta", {})
        rows.append(
            [
                str(probe.get("name", "")),
                format_float(vector_l2_norm(delta.get("eef_pos_delta", []))),
                format_float(vector_l2_norm(delta.get("gripper_qpos_delta", []))),
                format_float(vector_l2_norm(delta.get("joint_pos_delta", []))),
                format_float(probe.get("reward", 0.0)),
                str(bool(probe.get("done", False))).lower(),
                str(bool(probe.get("success", False))).lower(),
            ]
        )
    return rows


def build_top_level_metrics_table(artifacts: dict[str, Any]) -> MetricTable:
    open_loop = artifacts["open_loop"]
    vulnerability = artifacts["vulnerability"]
    closed_loop = next(iter(artifacts["closed_loop"].values()))
    eval_metrics = artifacts["eval_metrics"]
    rows = [
        ["eval.action_mse", format_float(eval_metrics.get("action_mse"))],
        ["open_loop.raw_action_mae", format_float(open_loop["raw_action_mae"])],
        ["open_loop.raw_action_mse", format_float(open_loop["raw_action_mse"])],
        ["open_loop.first_step_mae", format_float(open_loop["first_step_mae"])],
        ["open_loop.non_gripper_mae", format_float(open_loop["non_gripper_mae"])],
        ["open_loop.gripper_accuracy", format_percent(float(open_loop["gripper_accuracy"]))],
        ["vulnerability.wrong_instruction_shift_l2", format_float(vulnerability["wrong_instruction_shift_l2"])],
        ["vulnerability.blank_image_shift_l2", format_float(vulnerability["blank_image_shift_l2"])],
        ["vulnerability.zero_state_shift_l2", format_float(vulnerability["zero_state_shift_l2"])],
        ["vulnerability.mean_norm_saturation_fraction", format_float(vulnerability["mean_norm_saturation_fraction"])],
        ["closed_loop.average_success_rate", format_percent(float(closed_loop["average_success_rate"]))],
    ]
    return MetricTable(
        title="附录 A：完整 top-level metrics 表",
        columns=["metric", "value"],
        rows=rows,
    )


def build_closed_loop_appendix_table(closed_loop: dict[str, Any]) -> MetricTable:
    suite = next(iter(closed_loop.values()))
    rows = []
    for task_name, task in suite["tasks"].items():
        rows.append(
            [
                str(task["task_id"]),
                task_name,
                task["instruction"],
                str(task["successes"]),
                str(task["episodes"]),
                format_percent(float(task["success_rate"])),
            ]
        )
    return MetricTable(
        title="附录 B：2x10 每 task success 表",
        columns=["task_id", "task", "instruction", "successes", "episodes", "success_rate"],
        rows=rows,
    )


def build_vulnerability_appendix_table(vulnerability: dict[str, Any]) -> MetricTable:
    rows = []
    for sample in build_probe_sample_rows(vulnerability):
        rows.append(
            [
                str(sample.frame_index),
                sample.task_text,
                format_float(sample.wrong_instruction_shift_l2),
                format_float(sample.blank_image_shift_l2),
                format_float(sample.zero_state_shift_l2),
                format_float(sample.base_first_step_mae),
                format_float(sample.norm_saturation_fraction),
            ]
        )
    return MetricTable(
        title="附录 C：vulnerability probe 的 40 条样本级标量表",
        columns=[
            "frame_index",
            "task_text",
            "wrong_instruction_shift_l2",
            "blank_image_shift_l2",
            "zero_state_shift_l2",
            "base_first_step_mae",
            "norm_saturation_fraction",
        ],
        rows=rows,
    )


def build_action_semantics_appendix_table(semantics: dict[str, Any]) -> MetricTable:
    return MetricTable(
        title="附录 D：action semantics probe 名称与观测 delta 表",
        columns=[
            "probe",
            "eef_pos_delta_l2",
            "gripper_qpos_delta_l2",
            "joint_pos_delta_l2",
            "reward",
            "done",
            "success",
        ],
        rows=build_action_semantics_rows(semantics),
    )


def build_runtime_contract_table(artifacts: dict[str, Any]) -> MetricTable:
    pre = artifacts["preprocessor"]
    post = artifacts["postprocessor"]
    runtime = post.get("runtime", {})
    bridge = runtime.get("bridge", {})
    rows = [
        ["policy_version", str(pre.get("version", ""))],
        ["processor_model_id", str(pre.get("processor", {}).get("model_id", ""))],
        ["image_order", " | ".join(pre.get("input", {}).get("images", []))],
        ["prompt_template", str(pre.get("prompt", {}).get("template", ""))],
        ["state_format", str(pre.get("state", {}).get("format", ""))],
        ["action_norm", str(pre.get("normalization", {}).get("action", ""))],
        ["state_norm", str(pre.get("normalization", {}).get("state", ""))],
        ["gripper_norm", str(pre.get("normalization", {}).get("gripper", ""))],
        ["runtime_bridge_type", str(bridge.get("type", ""))],
        ["tap_strategy", str(bridge.get("tap_strategy", ""))],
        ["stop_gradient_backbone", str(bridge.get("stop_gradient_backbone", "")).lower()],
        ["bridge_out_dim", str(bridge.get("bridge_out_dim", ""))],
        ["action_expert_variant", str(bridge.get("action_expert_variant", ""))],
        ["conditioning_mode", str(bridge.get("conditioning_mode", ""))],
        ["use_memory_summary_conditioning", str(bridge.get("use_memory_summary_conditioning", "")).lower()],
        ["use_text_summary_conditioning", str(bridge.get("use_text_summary_conditioning", "")).lower()],
        ["use_instruction_summary_conditioning", str(bridge.get("use_instruction_summary_conditioning", "")).lower()],
        ["use_state_conditioning", str(bridge.get("use_state_conditioning", "")).lower()],
        ["use_action_input_conditioning", str(bridge.get("use_action_input_conditioning", "")).lower()],
        ["action_dim", str(post.get("action_dim", ""))],
        ["action_horizon", str(post.get("action_horizon", ""))],
    ]
    return MetricTable(
        title="附录 E：当前 checkpoint / runtime contract 摘要表",
        columns=["field", "value"],
        rows=rows,
    )


def write_cover(document: Document, metadata: ReportMetadata) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Qwen3.5 v3.1 Instrpool 模型诊断报告")
    add_run_with_font(run, "Microsoft YaHei", 22, bold=True)

    add_paragraph_text(document, f"Checkpoint: {metadata.checkpoint_dir}", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_text(document, f"生成时间: {metadata.generation_date}", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_text(document, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_labeled_paragraph(document, "摘要", metadata.summary_judgment)
    add_labeled_paragraph(
        document,
        "事实",
        "本报告只覆盖最新 v3.1 instrpool checkpoint，不展开历史版本对比；所有量化指标均来自 checkpoint 内现有 artifact。",
    )
    add_labeled_paragraph(
        document,
        "推断",
        "当前模型最主要的问题已经从“指令几乎不进入 expert”转移到“条件进入 expert 的方式仍然比 pi0.5 粗”。",
    )
    document.add_page_break()


def write_toc(document: Document) -> None:
    document.add_heading("目录", level=1)
    for item in [
        "第一章 模型总览",
        "第二章 架构拆解与归类",
        "第三章 pi0.5 对齐诊断",
        "第四章 诊断参数与行为证据",
        "第五章 问题定位与我的判断",
        "第六章 下一步验证建议",
        "附录",
    ]:
        add_paragraph_text(document, item)
    document.add_page_break()


def write_chapter_one(document: Document, metadata: ReportMetadata, artifacts: dict[str, Any]) -> None:
    document.add_heading("第一章 模型总览", level=1)
    add_labeled_paragraph(
        document,
        "事实",
        "当前实际运行链路可以概括为：prompt/state -> Qwen prefix -> last full-attn tap -> bridge -> summaries/conditioning -> Gemma expert -> flow-matching sampler。",
    )
    add_labeled_paragraph(
        document,
        "事实",
        "最新 checkpoint 身份为最新 v3.1 instrpool 版本，architecture 为 qwen_3.5_2b_gemma_bridge_v3_1，action head 为 pi05_gemma_bridge。",
    )
    add_labeled_paragraph(
        document,
        "推断",
        "从功能上看，这一版已经是“保留完整 Qwen backbone，再让 Gemma action expert 读取其 memory”的稳定实现，而不是早期的单纯 cross-attn 试验件。",
    )
    add_table(
        document,
        MetricTable(
            title="第一章 参数总表",
            columns=["参数", "值"],
            rows=build_core_config_rows(artifacts["config"], artifacts["postprocessor"]),
        ),
    )


def write_chapter_two(document: Document) -> None:
    document.add_heading("第二章 架构拆解与归类", level=1)
    add_labeled_paragraph(
        document,
        "事实",
        "本章按模块把当前 v3.1 拆成可定位问题的结构单元，并明确区分哪些是 pi0.5-style 核心、哪些是为模拟 pi0.5 功能而增加的替代实现。",
    )
    rows = [
        [
            item.module,
            item.current_structure,
            item.is_pi05_core,
            item.imitation_method,
            item.why,
            item.risk,
        ]
        for item in build_architecture_rows()
    ]
    add_table(
        document,
        MetricTable(
            title="第二章 模块归类表",
            columns=[
                "模块",
                "当前 v3.1 的实际构成",
                "是否属于 pi0.5-style 核心",
                "如果不是，如何仿造 pi0.5",
                "为什么这样做",
                "潜在失真/风险",
            ],
            rows=rows,
        ),
        font_size=8,
    )


def write_chapter_three(document: Document, code_evidence: dict[str, bool]) -> None:
    document.add_heading("第三章 pi0.5 对齐诊断", level=1)
    document.add_heading("3.1 原本就是 pi0.5-style 的部分", level=2)
    add_labeled_paragraph(
        document,
        "事实",
        "当前 v3.1 仍保留了 chunked action prediction、flow matching、reverse-time Euler、独立 action expert、continuous state conditioning、per-layer AdaRMS/FiLM 这些 pi0.5-style 核心协议。",
    )
    add_labeled_paragraph(
        document,
        "风险",
        "虽然这些控制协议被保留下来，但它们并不足以自动保证与 pi0.5 等价；真正决定行为质量的，是条件在 token 流中的传播方式。",
    )

    document.add_heading("3.2 不是 pi0.5 原生，但为模拟 pi0.5 而仿造的部分", level=2)
    add_labeled_paragraph(
        document,
        "事实",
        "Qwen full-attn tap memory、显式 bridge、Gemma cross-attention expert、memory/text/instruction summary、action-token global conditioning 都不是 pi0.5 原生结构。",
    )
    add_labeled_paragraph(
        document,
        "推断",
        "这些部件的目标都很明确：在不能直接复刻 pi0.5 prefix/suffix shared-stack 的前提下，用 external memory 和分层调制去近似实现“前缀影响动作 expert”的能力。",
    )
    add_labeled_paragraph(
        document,
        "风险",
        "模仿的是功能，不是同一套 operator 和交互拓扑；因此“看起来像条件进入 expert”并不等同于“行为上像 pi0.5”。",
    )

    document.add_heading("3.3 纯工程性稳定/诊断件", level=2)
    add_labeled_paragraph(
        document,
        "事实",
        "freeze Qwen、stop-gradient backbone、bridge gate bias、summary gain init、memory norm ratio limit、text_attention_mask 与 instruction pooled summary 的拆分，都属于稳定性和诊断辅助件。",
    )
    add_labeled_paragraph(
        document,
        "推断",
        "这些工程件成功解决了“条件几乎不可见”的早期问题，但它们更像补偿件，不是能把 Gemma bridge 自动变成 pi0.5-style token-level coupling 的核心机制。",
    )
    add_labeled_paragraph(
        document,
        "事实",
        "静态代码检查结果：adapter API 完整=%s，Gemma cross-attention 结构存在=%s，pi0.5 dual-stream 参考存在=%s。"
        % (
            format_bool(code_evidence["adapter_api_present"]),
            format_bool(code_evidence["gemma_cross_attention_present"]),
            format_bool(code_evidence["pi05_dual_stream_present"]),
        ),
    )


def write_chapter_four(document: Document, metadata: ReportMetadata, artifacts: dict[str, Any], code_evidence: dict[str, bool]) -> None:
    document.add_heading("第四章 诊断参数与行为证据", level=1)
    add_table(document, build_structure_metric_table(metadata, artifacts, code_evidence))
    add_labeled_paragraph(
        document,
        "解释",
        "这一组证据说明当前 checkpoint 是一个结构上自洽、配置上明确冻结 backbone 的 smoke1000 成果物。它能稳定保存并被后续脚本消费，但单测运行时标量没有落盘，所以这里的结论必须保持在 artifact 可证范围内。",
    )

    add_table(document, build_open_loop_table(artifacts["open_loop"]))
    add_labeled_paragraph(
        document,
        "解释",
        "open-loop 指标说明模型已经学到了一定程度的动作分布与 gripper 通道行为，并非完全随机或数值退化；但 first-step 和 non-gripper 指标仍然离闭环可用策略有明显距离。",
    )

    add_table(document, build_vulnerability_table(artifacts["vulnerability"]))
    add_labeled_paragraph(
        document,
        "解释",
        "vulnerability probe 是当前最关键的行为证据：wrong_instruction_shift 和 zero_state_shift 已经达到可见量级，说明文本与状态都能影响 expert；同时 blank_image_shift 更高，说明视觉仍然是更强的主导条件源。",
    )

    add_table(document, build_closed_loop_table(artifacts["closed_loop"]))
    add_labeled_paragraph(
        document,
        "解释",
        "closed-loop 2x10 子集说明运行时稳定性已经不是主要问题，但策略质量仍弱。换言之，现在不是“跑不起来”，而是“跑得起来但还不像 pi0.5 那样能稳定把条件落实成可执行轨迹”。",
    )


def write_chapter_five(document: Document, artifacts: dict[str, Any], code_evidence: dict[str, bool]) -> None:
    document.add_heading("第五章 问题定位与我的判断", level=1)
    document.add_heading("5.1 测得事实", level=2)
    for finding in build_findings(artifacts, code_evidence):
        if finding.kind != "事实":
            continue
        add_labeled_paragraph(document, finding.kind, finding.detail)

    document.add_heading("5.2 我的推断", level=2)
    for finding in build_findings(artifacts, code_evidence):
        if finding.kind != "推断":
            continue
        add_labeled_paragraph(document, finding.kind, finding.detail)

    document.add_heading("5.3 最可疑的主因排序", level=2)
    add_table(
        document,
        MetricTable(
            title="第五章 主因排序表",
            columns=["排名", "主因假设", "为什么更像这个问题", "为什么暂时不像其他问题", "结论"],
            rows=build_suspicion_rows(),
        ),
        font_size=8,
    )


def write_chapter_six(document: Document) -> None:
    document.add_heading("第六章 下一步验证建议", level=1)
    for item in build_next_steps():
        add_paragraph_text(document, item)


def write_appendix(document: Document, artifacts: dict[str, Any]) -> None:
    document.add_heading("附录", level=1)
    add_table(document, build_top_level_metrics_table(artifacts))
    add_table(document, build_closed_loop_appendix_table(artifacts["closed_loop"]), font_size=8)
    add_table(document, build_vulnerability_appendix_table(artifacts["vulnerability"]), font_size=8)
    add_table(document, build_action_semantics_appendix_table(artifacts["semantics"]))
    add_table(document, build_runtime_contract_table(artifacts), font_size=8)


def generate_document(metadata: ReportMetadata, artifacts: dict[str, Any], code_evidence: dict[str, bool]) -> Document:
    document = Document()
    set_document_defaults(document)
    write_cover(document, metadata)
    write_toc(document)
    write_chapter_one(document, metadata, artifacts)
    write_chapter_two(document)
    write_chapter_three(document, code_evidence)
    write_chapter_four(document, metadata, artifacts, code_evidence)
    write_chapter_five(document, artifacts, code_evidence)
    write_chapter_six(document)
    write_appendix(document, artifacts)
    return document


def main() -> None:
    args = parse_args()
    checkpoint_dir = args.checkpoint_dir.resolve()
    output_path = args.output.resolve() if args.output else checkpoint_dir / "Qwen35_v31_instrpool_diagnostic_report_zh.docx"
    repo_root = Path(__file__).resolve().parent

    artifacts = load_artifacts(checkpoint_dir)
    code_evidence = collect_code_evidence(repo_root)
    metadata = build_metadata(checkpoint_dir, output_path, args.language, artifacts)
    document = generate_document(metadata, artifacts, code_evidence)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    print(json.dumps({"output": str(output_path), "checkpoint_dir": str(checkpoint_dir)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
